"""업무 문서(Word/PDF/엑셀/텍스트)에서 LLM에 보낼 텍스트를 추출."""

from __future__ import annotations

import csv
import re
import zipfile
from pathlib import Path

SUPPORTED_SUFFIXES = {".docx", ".pdf", ".xlsx", ".xlsm", ".csv", ".txt", ".md"}


def find_documents(root: Path) -> list[Path]:
    """root 아래의 지원 문서를 정렬해서 반환 (임시 파일 ~$ 제외)."""
    return sorted(
        p
        for p in root.rglob("*")
        if p.is_file()
        and p.suffix.lower() in SUPPORTED_SUFFIXES
        and not p.name.startswith(("~$", "."))
    )


_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # 옛 Office(97-2003)·암호화 Office 파일의 공통 헤더
_KIND = {".docx": ("Word", "Word 문서(.docx)"), ".xlsx": ("엑셀", "Excel 통합 문서(.xlsx)"), ".xlsm": ("엑셀", "Excel 통합 문서(.xlsx)")}


def _utf16(name: str) -> bytes:
    return name.encode("utf-16-le")


def check_format(name: str, data: bytes) -> None:
    """확장자와 실제 파일 내용이 맞는지 확인. 읽을 수 없는 파일이면 이유를 담은 ValueError.

    회사 보안 프로그램(DRM)으로 암호화된 파일, 비밀번호가 걸린 파일,
    옛 형식(.doc/.xls)에 확장자만 바꾼 파일을 업로드 단계에서 알아듣기 쉽게 알려준다.
    """
    suffix = Path(name).suffix.lower()
    if not data:
        raise ValueError("빈 파일입니다.")
    if suffix in _KIND:
        app, save_as = _KIND[suffix]
        if data.startswith(b"PK"):
            return
        if data.startswith(_OLE_MAGIC):
            if _utf16("EncryptionInfo") in data or _utf16("EncryptedPackage") in data:
                raise ValueError(
                    f"비밀번호가 걸린 {app} 파일이라 읽을 수 없습니다. "
                    f"{app}에서 [파일 → 정보 → 문서 보호]로 암호를 해제하고 저장한 뒤 다시 올려 주세요."
                )
            raise ValueError(
                f"옛 {app} 97-2003 형식(.doc/.xls) 파일입니다. "
                f"{app}에서 [다른 이름으로 저장 → {save_as}]로 저장한 뒤 다시 올려 주세요."
            )
        if data.lstrip()[:5] == b"%PDF-":
            raise ValueError("내용은 PDF인데 확장자가 다릅니다. 확장자를 .pdf로 바꿔서 올려 주세요.")
        raise ValueError(
            f"{app} 파일 형식이 아닙니다. 회사 보안 프로그램(DRM)으로 암호화된 파일일 수 있습니다. "
            f"보안을 해제하거나, {app}에서 PDF로 저장해서 올려 주세요."
        )
    if suffix == ".pdf" and b"%PDF-" not in data[:1024]:
        raise ValueError(
            "PDF 파일 형식이 아닙니다. 회사 보안 프로그램(DRM)으로 암호화된 파일일 수 있습니다. "
            "보안을 해제한 뒤 다시 올려 주세요."
        )


def _docx_xml_text(path: Path) -> str:
    """python-docx가 못 여는 비표준 .docx를 위한 대체 추출: 본문 XML에서 글자만 뽑는다."""
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="replace")
    xml = re.sub(r"</w:p>|<w:br[^>]*/>|<w:tab[^>]*/>", "\n", xml)
    text = re.sub(r"<[^>]+>", "", xml)
    for ent, ch in (("&lt;", "<"), ("&gt;", ">"), ("&quot;", '"'), ("&apos;", "'"), ("&amp;", "&")):
        text = text.replace(ent, ch)
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def _docx_text(path: Path) -> str:
    import docx  # python-docx

    try:
        d = docx.Document(str(path))
    except Exception:
        # 형식 문제면 이유를 알려주고, 구조만 조금 다른 정상 파일이면 XML에서 직접 읽는다.
        check_format(path.name, path.read_bytes())
        try:
            return _docx_xml_text(path)
        except (zipfile.BadZipFile, KeyError) as e:
            raise ValueError("Word 문서를 읽지 못했습니다. Word에서 다시 저장하거나 PDF로 저장해서 올려 주세요.") from e
    parts: list[str] = [p.text for p in d.paragraphs if p.text.strip()]
    for ti, table in enumerate(d.tables, 1):
        parts.append(f"\n[표 {ti}]")
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx_text(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"\n[시트: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _csv_text(path: Path) -> str:
    for enc in ("utf-8-sig", "cp949"):
        try:
            with path.open(encoding=enc, newline="") as f:
                return "\n".join(" | ".join(r) for r in csv.reader(f))
        except UnicodeDecodeError:
            continue
    raise ValueError(f"CSV 인코딩을 읽을 수 없습니다: {path}")


def _plain_text(path: Path) -> str:
    for enc in ("utf-8-sig", "cp949"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"텍스트 인코딩을 읽을 수 없습니다: {path}")


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("비밀번호가 걸린 PDF라 읽을 수 없습니다. 암호를 해제하고 다시 저장해서 올려 주세요.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError as e:
        raise ValueError("PDF를 읽지 못했습니다. 파일이 손상되었거나 보안이 걸려 있을 수 있습니다.") from e
    if not any(p.strip() for p in pages):
        raise ValueError(
            "PDF에 글자 정보가 없습니다(스캔한 이미지 PDF). 원본 Word/한글 파일을 올리거나 OCR 처리한 PDF를 올려 주세요."
        )
    return "\n\n".join(f"[{i}쪽]\n{t.strip()}" for i, t in enumerate(pages, 1) if t.strip())


def extract_text(path: Path) -> str:
    """문서 하나에서 LLM에 보낼 텍스트를 뽑는다. 읽을 수 없으면 이유를 담은 ValueError."""
    suffix = path.suffix.lower()
    if suffix in (".pdf", ".xlsx", ".xlsm"):
        check_format(path.name, path.read_bytes())
    if suffix == ".pdf":
        return _pdf_text(path)
    if suffix == ".docx":
        return _docx_text(path)
    if suffix in (".xlsx", ".xlsm"):
        return _xlsx_text(path)
    if suffix == ".csv":
        return _csv_text(path)
    if suffix in (".txt", ".md"):
        return _plain_text(path)
    raise ValueError(f"지원하지 않는 형식: {path.suffix}")
