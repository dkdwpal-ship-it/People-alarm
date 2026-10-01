import docx
import openpyxl

from people_alarm.loaders import extract_text, find_documents


def test_docx_and_xlsx_to_text(tmp_path):
    d = docx.Document()
    d.add_paragraph("매월 10일까지 원천세를 신고한다.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "연말정산"
    table.rows[0].cells[1].text = "1월"
    d.save(tmp_path / "manual.docx")

    wb = openpyxl.Workbook()
    wb.active.title = "연간일정"
    wb.active.append(["월", "업무"])
    wb.active.append([3, "법인세 신고"])
    wb.save(tmp_path / "calendar.xlsx")

    (tmp_path / "~$temp.docx").write_bytes(b"lock")
    (tmp_path / "image.png").write_bytes(b"x")

    docs = find_documents(tmp_path)
    assert [p.name for p in docs] == ["calendar.xlsx", "manual.docx"]

    xlsx_text = extract_text(docs[0])
    assert "[시트: 연간일정]" in xlsx_text and "3 | 법인세 신고" in xlsx_text

    docx_text = extract_text(docs[1])
    assert "원천세" in docx_text and "연말정산 | 1월" in docx_text


def _make_pdf(text: str | None) -> bytes:
    """글자 한 줄짜리(또는 글자 없는) 최소 PDF."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def test_pdf_text_extraction(tmp_path):
    (tmp_path / "a.pdf").write_bytes(_make_pdf("Payroll on the 25th"))
    assert "Payroll on the 25th" in extract_text(tmp_path / "a.pdf")


def test_scanned_pdf_without_text(tmp_path):
    (tmp_path / "scan.pdf").write_bytes(_make_pdf(None))
    with pytest.raises(ValueError, match="스캔"):
        extract_text(tmp_path / "scan.pdf")


import zipfile

import pytest

from people_alarm.loaders import check_format

OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


@pytest.mark.parametrize("name, data, expected", [
    ("a.docx", OLE + b"\0" * 64 + "EncryptionInfo".encode("utf-16-le"), "비밀번호"),
    ("a.docx", OLE + b"\0" * 64 + "WordDocument".encode("utf-16-le"), "97-2003"),
    ("a.xlsx", OLE + b"\0" * 64 + "Workbook".encode("utf-16-le"), "Excel 통합 문서"),
    ("a.docx", b"<## DRM encrypted ##>" + b"\0" * 32, "DRM"),
    ("a.docx", b"%PDF-1.7 ...", "확장자를 .pdf"),
    ("a.pdf", b"\x00\x01garbage", "DRM"),
    ("a.docx", b"", "빈 파일"),
])
def test_check_format_explains_unreadable_files(name, data, expected):
    with pytest.raises(ValueError, match=expected):
        check_format(name, data)


def test_check_format_accepts_real_files():
    check_format("a.docx", b"PK\x03\x04rest")
    check_format("a.pdf", b"%PDF-1.4 rest")
    check_format("a.txt", b"anything")


def test_docx_fallback_for_nonstandard_package(tmp_path):
    # python-docx가 요구하는 [Content_Types].xml 없이 본문만 있는 비표준 docx
    path = tmp_path / "odd.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml",
                   '<w:document><w:body><w:p><w:r><w:t>매월 10일 원천세 신고 &amp; 납부</w:t></w:r></w:p>'
                   '<w:p><w:r><w:t>분기 말 성과 점검</w:t></w:r></w:p></w:body></w:document>')
    text = extract_text(path)
    assert "매월 10일 원천세 신고 & 납부\n분기 말 성과 점검" in text


def test_unreadable_docx_gives_friendly_error(tmp_path):
    path = tmp_path / "drm.docx"
    path.write_bytes(b"<## DRM ##>" + b"\0" * 100)
    with pytest.raises(ValueError, match="DRM"):
        extract_text(path)
